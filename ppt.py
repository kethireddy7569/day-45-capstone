from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BLUE = RGBColor(0, 95, 180)
DARK = RGBColor(30, 45, 65)
LIGHT = RGBColor(235, 244, 252)
GRAY = RGBColor(90, 100, 110)
WHITE = RGBColor(255, 255, 255)


def add_text(slide, x, y, w, h, text, size=20,
             color=DARK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    tf = box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = "Arial"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def new_slide(title, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    add_text(slide, 0.7, 0.35, 11.8, 0.6,
             title, 28, BLUE, True)

    add_text(slide, 12.1, 7.05, 0.5, 0.25,
             str(number), 10, GRAY, False, PP_ALIGN.RIGHT)

    return slide


def add_bullets(slide, items, y=1.5, size=20):
    box = slide.shapes.add_textbox(
        Inches(0.9), Inches(y), Inches(11.5), Inches(5.3)
    )
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True

    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + item
        p.font.name = "Arial"
        p.font.size = Pt(size)
        p.font.color.rgb = DARK
        p.space_after = Pt(16)


def add_card(slide, x, y, w, h, heading, description):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = RGBColor(205, 225, 245)

    add_text(slide, x + 0.18, y + 0.15, w - 0.36, 0.4,
             heading, 17, BLUE, True)

    add_text(slide, x + 0.18, y + 0.65, w - 0.36, h - 0.75,
             description, 14, DARK)


# SLIDE 1: TITLE PAGE
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = WHITE

add_text(slide, 2.2, 1.2, 9, 0.8,
         "CRAFTECHCO", 44, BLUE, True, PP_ALIGN.CENTER)

add_text(slide, 2.2, 2.0, 9, 0.5,
         "IT SOLUTIONS PVT. LTD.", 22, BLUE,
         False, PP_ALIGN.CENTER)

add_text(slide, 1.0, 4.0, 11.3, 0.6,
         "ENTERPRISE GENAI SYSTEMS ENGINEER",
         25, DARK, True, PP_ALIGN.CENTER)

add_text(slide, 1.0, 4.7, 11.3, 0.4,
         "Capstone Project Presentation | Day 1–45",
         18, GRAY, False, PP_ALIGN.CENTER)

add_text(slide, 1.0, 5.7, 11.3, 0.4,
         "Presented by: Akhila Kethireddy",
         20, BLUE, True, PP_ALIGN.CENTER)

add_text(slide, 1.0, 6.2, 11.3, 0.4,
         "Employee ID: CRFT-IT-260601",
         17, DARK, False, PP_ALIGN.CENTER)


# SLIDE 2: PROJECT OVERVIEW
slide = new_slide("Project Overview", 2)

add_bullets(slide, [
    "This project summarizes the 45-day Enterprise GenAI Systems Engineer training.",
    "The training covered Python scripting, API integration, AI workflows and data handling.",
    "Different task modules were developed to understand practical software systems.",
    "The capstone presents the overall project architecture and workflow."
])


# SLIDE 3: OBJECTIVES
slide = new_slide("Project Objectives", 3)

add_card(slide, 0.8, 1.5, 3.7, 1.7,
         "Automation", "Automate repetitive tasks using Python scripts.")

add_card(slide, 4.8, 1.5, 3.7, 1.7,
         "Integration", "Connect APIs, AI services and databases.")

add_card(slide, 8.8, 1.5, 3.7, 1.7,
         "Monitoring", "Capture relevant activity and system information.")

add_card(slide, 2.8, 3.8, 3.7, 1.7,
         "Validation", "Validate inputs and handle errors.")

add_card(slide, 6.8, 3.8, 3.7, 1.7,
         "Documentation", "Record implementation and testing details.")


# SLIDE 4: TECHNOLOGY STACK
slide = new_slide("Technology Stack", 4)

add_card(slide, 0.8, 1.5, 3.7, 1.7,
         "Programming", "Python, HTML, CSS and JavaScript")

add_card(slide, 4.8, 1.5, 3.7, 1.7,
         "AI Integration", "OpenRouter and Groq")

add_card(slide, 8.8, 1.5, 3.7, 1.7,
         "Database", "Firebase and Firestore")

add_card(slide, 0.8, 3.8, 3.7, 1.7,
         "Asset Management", "Snipe-IT")

add_card(slide, 4.8, 3.8, 3.7, 1.7,
         "Web Data", "Context.dev")

add_card(slide, 8.8, 3.8, 3.7, 1.7,
         "Development", "Git, GitHub and deployment platforms")


# SLIDE 5: SYSTEM ARCHITECTURE
slide = new_slide("System Architecture", 5)

architecture = [
    ("User Input", 0.7),
    ("Python Backend", 3.2),
    ("AI / APIs", 5.7),
    ("Database", 8.2),
    ("Dashboard", 10.7)
]

for label, x in architecture:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(2.4), Inches(1.9), Inches(1.0)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = BLUE

    add_text(slide, x + 0.08, 2.65, 1.74, 0.4,
             label, 14, DARK, True, PP_ALIGN.CENTER)

for x in [2.65, 5.15, 7.65, 10.15]:
    arrow = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW,
        Inches(x), Inches(2.7), Inches(0.45), Inches(0.3)
    )
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = BLUE
    arrow.line.fill.background()

add_text(slide, 1.0, 4.4, 11.3, 0.8,
         "Input → Processing → API Integration → Data Storage → Dashboard",
         20, DARK, True, PP_ALIGN.CENTER)


# SLIDE 6: WORKFLOW
slide = new_slide("End-to-End Workflow", 6)

add_bullets(slide, [
    "Step 1: Receive user input or task data.",
    "Step 2: Validate required fields and data formats.",
    "Step 3: Process information using Python.",
    "Step 4: Call the required AI service or external API.",
    "Step 5: Store results and logs in the configured data store.",
    "Step 6: Display the results through the interface or dashboard."
], size=19)


# SLIDE 7: MAIN WORK AREAS
slide = new_slide("Main Work Areas", 7)

add_bullets(slide, [
    "Python scripting and REST API integration.",
    "Input validation and duplicate record handling.",
    "AI-assisted help desk and question-answering workflows.",
    "Job data collection and filtering.",
    "Activity tracking and dashboard development.",
    "Frontend development and deployment troubleshooting."
], size=19)


# SLIDE 8: SECURITY AND RELIABILITY
slide = new_slide("Security and Reliability", 8)

add_bullets(slide, [
    "Store API keys in environment variables or secure platform settings.",
    "Never commit .env files or service account credentials to GitHub.",
    "Validate API responses and handle errors clearly.",
    "Check database permissions and required cloud APIs.",
    "Review deployment logs when a service fails.",
    "Test each module before integrating it into the overall workflow."
], size=19)


# SLIDE 9: TESTING AND VALIDATION
slide = new_slide("Testing and Validation", 9)

add_bullets(slide, [
    "Input Testing: Test valid and invalid sample data.",
    "API Testing: Verify status codes and response content.",
    "Database Testing: Confirm that expected records are saved.",
    "Dashboard Testing: Check that results are displayed correctly.",
    "Error Testing: Confirm that failures are handled safely.",
    "Deployment Testing: Verify environment configuration and logs."
], size=19)


# SLIDE 10: CONCLUSION
slide = new_slide("Conclusion and Learning Outcomes", 10)

add_bullets(slide, [
    "Improved understanding of Python automation and API integration.",
    "Learned how AI services and data stores can work together.",
    "Practiced validation, logging, error handling and deployment.",
    "Developed troubleshooting and technical documentation skills.",
    "The project provides a foundation for further system development."
], size=19)

add_text(slide, 1.0, 6.25, 11.3, 0.45,
         "THANK YOU", 25, BLUE, True, PP_ALIGN.CENTER)


# Save the presentation
prs.save("Akhila_Kethireddy_Capstone_Presentation.pptx")
print("PowerPoint created successfully!")